import json
import logging
import os
from pathlib import Path
from typing import Any

import azure.functions as func
import azure.durable_functions as df

from agent_summary import summarize_report
from dts_client import normalize_with_dts


app = func.FunctionApp(http_auth_level=func.AuthLevel.FUNCTION)


@app.route(route="etl/start", methods=["POST"])
@app.durable_client_input(client_name="client")
async def start_etl(req: func.HttpRequest, client: df.DurableOrchestrationClient):
    payload = req.get_json() if req.get_body() else {}
    input_folder = payload.get("input_folder", os.getenv("ETL_INPUT_FOLDER", "data/incoming"))
    instance_id = await client.start_new(
        "etl_orchestrator",
        client_input={"input_folder": input_folder},
    )
    status_url = f"{req.url.rsplit('/api/', 1)[0]}/api/etl/status/{instance_id}"
    return func.HttpResponse(
        json.dumps({"id": instance_id, "statusQueryGetUri": status_url}),
        status_code=202,
        mimetype="application/json",
    )


@app.route(route="etl/status/{instance_id}", methods=["GET"])
@app.durable_client_input(client_name="client")
async def get_etl_status(req: func.HttpRequest, client: df.DurableOrchestrationClient):
    instance_id = req.route_params["instance_id"]
    status = await client.get_status(instance_id)
    if status is None:
        return func.HttpResponse(
            json.dumps({"error": "Instance not found", "instance_id": instance_id}),
            status_code=404,
            mimetype="application/json",
        )
    return func.HttpResponse(
        json.dumps(
            {
                "name": status.name,
                "instanceId": status.instance_id,
                "runtimeStatus": status.runtime_status.name if status.runtime_status else None,
                "input": status.input_,
                "output": status.output,
                "createdTime": status.created_time.isoformat() if status.created_time else None,
                "lastUpdatedTime": status.last_updated_time.isoformat()
                if status.last_updated_time
                else None,
            }
        ),
        status_code=200,
        mimetype="application/json",
    )


@app.route(route="etl/cancel/{instance_id}", methods=["POST"])
@app.durable_client_input(client_name="client")
async def cancel_etl(req: func.HttpRequest, client: df.DurableOrchestrationClient):
    instance_id = req.route_params["instance_id"]
    input_folder = req.params.get("input_folder", os.getenv("ETL_INPUT_FOLDER", "data/incoming"))
    await client.terminate(instance_id, "Cancelled by user")
    cleanup_id = await client.start_new(
        "cleanup_orchestrator",
        client_input={"input_folder": input_folder},
    )
    return func.HttpResponse(
        json.dumps({"instance_id": instance_id, "cleanup_instance_id": cleanup_id}),
        status_code=202,
        mimetype="application/json",
    )


@app.orchestration_trigger(context_name="context")
def etl_orchestrator(context: df.DurableOrchestrationContext):
    request = context.get_input() or {}
    files = yield context.call_activity("discover_excel_files", request)
    summaries = yield context.task_all(
        [context.call_activity("normalize_excel_file", file_path) for file_path in files]
    )
    report = yield context.call_activity(
        "create_word_report",
        {"input_folder": request.get("input_folder"), "summaries": summaries},
    )
    return report


@app.orchestration_trigger(context_name="context")
def cleanup_orchestrator(context: df.DurableOrchestrationContext):
    yield context.call_activity("erase_etl_data", context.get_input() or {})
    return {"status": "erased"}


@app.activity_trigger(input_name="request")
def discover_excel_files(request: dict) -> list[str]:
    input_folder = Path(request["input_folder"]).resolve()
    files = sorted(path for path in input_folder.glob("*.xls*") if path.is_file())
    if len(files) != 2:
        raise ValueError(f"Expected exactly two Excel files in {input_folder}, found {len(files)}")
    return [str(path) for path in files]


@app.activity_trigger(input_name="file_path")
def normalize_excel_file(file_path: str) -> dict[str, Any]:
    logging.info("Normalizing %s through local DTS", file_path)
    return normalize_with_dts(file_path)


@app.activity_trigger(input_name="request")
async def create_word_report(request: dict) -> dict[str, str]:
    input_folder = Path(request["input_folder"]).resolve()
    output_folder = Path(os.getenv("ETL_OUTPUT_FOLDER", input_folder.parent / "output")).resolve()
    output_folder.mkdir(parents=True, exist_ok=True)
    report_path = output_folder / "etl-report.docx"
    narrative = await summarize_report(request["summaries"])

    from docx import Document

    document = Document()
    document.add_heading("ETL Summary Report", level=1)
    document.add_paragraph(narrative)
    for summary in request["summaries"]:
        document.add_heading(summary["file_name"], level=2)
        document.add_paragraph(
            f"Rows: {summary['row_count']} | Columns: {', '.join(summary['columns'])}"
        )
        document.add_paragraph(f"Null cells: {summary['null_cells']}")
        if summary["numeric_totals"]:
            document.add_paragraph(json.dumps(summary["numeric_totals"], indent=2))
    document.save(report_path)
    return {"report_path": str(report_path), "status": "completed"}


@app.activity_trigger(input_name="request")
def erase_etl_data(request: dict) -> dict[str, Any]:
    input_folder = Path(request["input_folder"]).resolve()
    output_folder = Path(os.getenv("ETL_OUTPUT_FOLDER", input_folder.parent / "output")).resolve()
    work_folder = Path(os.getenv("ETL_WORK_FOLDER", input_folder.parent / "work")).resolve()
    for folder in (input_folder, output_folder, work_folder):
        if folder.exists():
            for child in folder.iterdir():
                if child.is_file() or child.is_symlink():
                    child.unlink()
                elif child.is_dir():
                    import shutil

                    shutil.rmtree(child)
    return {"folders": [str(input_folder), str(output_folder), str(work_folder)]}